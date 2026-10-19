"""Карта сети и паспорт актива (SDD §5.10): граф, связи, фильтры, выгрузка."""

from __future__ import annotations

from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.discovery.models import RouterDevice
from apps.discovery.tasks import scan_segment
from apps.registry import mapdata
from apps.registry.models import Asset, Segment, Site, System, SystemComponent

pytestmark = pytest.mark.django_db


@pytest.fixture
def site() -> Site:
    return Site.objects.create(name="Площадка-Карта")


@pytest.fixture
def segment(site: Site) -> Segment:
    return Segment.objects.create(
        site=site, name="Сегмент-Карта", cidr="10.70.1.0/24", vlan=70, scan_enabled=True
    )


@pytest.fixture
def router() -> RouterDevice:
    return RouterDevice.objects.create(
        host="10.0.0.1",
        name="MT-Demo",
        addresses=[{"address": "10.70.1.1/24", "interface": "bridge", "network": "10.70.1.0/24"}],
        last_seen_at=timezone.now(),
    )


def make_asset(segment: Segment, ip: str, **fields) -> Asset:
    return Asset.objects.create(
        ip=ip, segment=segment, site=segment.site, last_seen_at=timezone.now(), **fields
    )


def test_overview_graph_has_site_and_router_segment_edge(
    site: Site, segment: Segment, router: RouterDevice
) -> None:
    graph = mapdata.overview_graph()

    node_ids = {node["data"]["id"] for node in graph["nodes"]}
    assert f"site:{site.pk}" in node_ids
    assert f"segment:{segment.pk}" in node_ids
    assert f"router:{router.pk}" in node_ids
    assert any(
        edge["data"]["source"] == f"router:{router.pk}"
        and edge["data"]["target"] == f"segment:{segment.pk}"
        for edge in graph["edges"]
    )


def test_segment_graph_links_assets_to_systems(segment: Segment) -> None:
    system = System.objects.create(name="ЛИС-Карта", site=segment.site)
    linked = make_asset(
        segment, "10.70.1.10", hostname="lis-db", role_auto="server", mac_vendor="Dell Inc."
    )
    SystemComponent.objects.create(system=system, asset=linked)
    make_asset(segment, "10.70.1.11", hostname="unknown-1")

    graph = mapdata.segment_graph(segment, list(mapdata.asset_queryset().filter(segment=segment)))

    node_ids = {node["data"]["id"] for node in graph["nodes"]}
    assert f"asset:{linked.pk}" in node_ids
    assert f"system:{system.pk}" in node_ids
    assert any(
        edge["data"]["source"] == f"asset:{linked.pk}"
        and edge["data"]["target"] == f"system:{system.pk}"
        for edge in graph["edges"]
    )
    unknown = Asset.objects.get(hostname="unknown-1")
    assert not any(edge["data"]["source"] == f"asset:{unknown.pk}" for edge in graph["edges"])


def test_map_view_renders_overview_and_segment(login, make_user, segment: Segment) -> None:
    client = login(make_user("map-officer", role="ib_officer"))

    assert client.get("/registry/map/").status_code == 200
    assert client.get(f"/registry/map/?segment={segment.pk}").status_code == 200


def test_map_filters_by_role_vendor_and_status(login, make_user, segment: Segment) -> None:
    make_asset(segment, "10.70.1.20", role_auto="printer", mac_vendor="HP")
    make_asset(segment, "10.70.1.21", role_auto="server", mac_vendor="Dell")
    client = login(make_user("map-filter", role="ib_officer"))

    response = client.get(f"/registry/map/?segment={segment.pk}&role=printer")
    assert [asset.ip for asset in response.context["assets"]] == ["10.70.1.20"]

    response = client.get(f"/registry/map/?segment={segment.pk}&vendor=Dell")
    assert [asset.ip for asset in response.context["assets"]] == ["10.70.1.21"]

    Asset.objects.filter(ip="10.70.1.21").update(miss_count=3)
    response = client.get(f"/registry/map/?segment={segment.pk}&status=quiet")
    assert [asset.ip for asset in response.context["assets"]] == ["10.70.1.21"]


def test_map_sorting_by_ip(login, make_user, segment: Segment) -> None:
    make_asset(segment, "10.70.1.20")
    make_asset(segment, "10.70.1.5")
    client = login(make_user("map-sort", role="ib_officer"))

    response = client.get(f"/registry/map/?segment={segment.pk}&sort=-ip")

    assert [asset.ip for asset in response.context["assets"]] == ["10.70.1.20", "10.70.1.5"]


def test_segment_assets_xlsx_export(login, make_user, segment: Segment) -> None:
    make_asset(segment, "10.70.1.30", role_auto="printer", software="HP LaserJet 4")
    client = login(make_user("map-xlsx", role="ib_officer"))

    response = client.get(f"/registry/map/assets.xlsx?segment={segment.pk}")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    header = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
    assert header[:4] == ["IP", "MAC", "Вендор", "Имя"]
    assert workbook.active.max_row == 2


def test_segment_xlsx_requires_segment(login, make_user) -> None:
    client = login(make_user("map-xlsx-no-segment", role="ib_officer"))

    assert client.get("/registry/map/assets.xlsx").status_code == 404


def test_scan_now_enqueues_manual_scan(login, make_user, segment: Segment, monkeypatch) -> None:
    calls: dict = {}
    monkeypatch.setattr(scan_segment, "apply_async", lambda **kwargs: calls.update(kwargs))
    client = login(make_user("map-scan-officer", role="ib_officer"))

    response = client.post(f"/registry/segments/{segment.pk}/scan-now/")

    assert response.status_code == 302
    assert calls["args"] == [segment.pk, "manual"]
    assert calls["queue"] == "scan"


def test_scan_now_forbidden_for_reader(login, make_user, segment: Segment) -> None:
    client = login(make_user("map-scan-reader", role="reader"))

    assert client.post(f"/registry/segments/{segment.pk}/scan-now/").status_code == 403
