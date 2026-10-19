"""Неделя 11: отчёты — доступ, состав, выгрузки."""

from __future__ import annotations

import pytest

pytestmark = pytest.mark.django_db


def test_reports_requires_login(client) -> None:
    response = client.get("/reports/")

    assert response.status_code == 302
    assert "/accounts/login/" in response["Location"]


def test_reports_shows_sections(client, login, make_user) -> None:
    client = login(make_user("officer-rep", role="ib_officer"))

    body = client.get("/reports/").content.decode()

    assert "Состояние контролей" in body
    assert "Просроченные обязательства" in body
    assert "Покрытие evidence по видам" in body
    assert "Оценки КЗИ/ПЗИ" in body


def test_gaps_csv_headers_and_access(client, login, make_user) -> None:
    client = login(make_user("officer-csv", role="ib_officer"))

    response = client.get("/reports/gaps.csv")

    assert response.status_code == 200
    body = response.content.decode("utf-8-sig")
    assert body.startswith("Код;Контроль;Домен;ИС;Статус;Ответственный;Следующая проверка")


def test_reader_can_view_reports_and_export(client, login, make_user) -> None:
    """«Руководство» читает отчёты (в матрице ролей — только чтение)."""
    client = login(make_user("reader-rep", role="reader"))

    assert client.get("/reports/").status_code == 200
    assert client.get("/reports/overdue.csv").status_code == 200


def test_xlsx_exports_are_loadable(client, login, make_user) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    from apps.core.export import XLSX_CONTENT_TYPE

    client = login(make_user("officer-xlsx", role="ib_officer"))

    response = client.get("/reports/gaps.xlsx")

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX_CONTENT_TYPE
    workbook = load_workbook(BytesIO(response.content))
    header = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
    assert header[0] == "Код"
    for url in (
        "/reports/overdue.xlsx",
        "/reports/evidence.xlsx",
        "/reports/unknown-assets.xlsx",
        "/reports/kzi.xlsx",
    ):
        assert client.get(url).status_code == 200


def test_management_report_renders_one_pager(client, login, make_user) -> None:
    client = login(make_user("boss-mgmt", role="reader"))

    response = client.get("/reports/management/")
    body = response.content.decode()

    assert response.status_code == 200
    for marker in ("Отчёт руководству", "Риски", "План мероприятий", "КЗИ"):
        assert marker in body
